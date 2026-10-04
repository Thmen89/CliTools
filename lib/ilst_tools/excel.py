from __future__ import annotations

import getpass
import importlib.util
import math
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd

SCALE = {"1": "", "k": ",", "mn": ",,", "bn": ",,,"}
SCENARIO_COLOURS = ["1F4E79", "C00000", "BF8F00", "548235", "7030A0", "7F7F7F"]


def available() -> bool:
    return importlib.util.find_spec("openpyxl") is not None


def folder(export_dir=None) -> Path:
    if export_dir:
        f = Path(export_dir).expanduser()
        if not f.is_dir():
            from .core import DataError
            raise DataError(f"Export folder not found: {f} - change it in Settings.")
        return f
    downloads = Path.home() / "Downloads"
    return downloads if downloads.is_dir() else Path.home()


@dataclass
class Sheet:
    title: str
    frame: pd.DataFrame
    kinds: dict = field(default_factory=dict)
    lines: list = field(default_factory=list)
    source: str = ""
    total: bool = False
    flag: list = field(default_factory=list)
    chart: dict | None = None
    tab: str = ""


def result_sheet(res, title: str, units: str = "mn") -> Sheet:
    from .core import LABELS
    name = LABELS[res.view.breakdown]
    rows = [{name: k, **dict(zip(res.columns, v))} for k, v in res.rows]
    rows.append({name: "Total", **dict(zip(res.columns, res.totals))})
    note = f"USD {units} · negative = requirement" + (" · Change = B − A" if res.view.compare else "")
    return Sheet(title, pd.DataFrame(rows), {c: "money" for c in res.columns}, list(res.context) + [note],
                 ", ".join(dict.fromkeys(res.sources)), total=True, tab="Summary")


def data_sheet(frame: pd.DataFrame, res) -> Sheet:
    kinds = {"Report": "date", "Raw USD": "money", "USD": "money", "Factor": "pct", "Day": "int"}
    lines = list(res.context[:2 if res.view.compare else 1]) + [
        "Underlying rows behind the summary, one per combination of dimensions - filter or pivot freely.",
        "USD = Raw USD x Factor (prepositioning on M2–M12 flows when applied); the USD column sums to the summary."]
    return Sheet("Data", frame, kinds, lines, ", ".join(dict.fromkeys(res.sources)), tab="Data")


def write(path: Path, sheets: list, units: str = "mn", decimals: int = 1) -> Path:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    base = "#,##0" + ("." + "0" * decimals if decimals else "") + SCALE.get(units, ",,")
    fmt = {"money": f'{base};-{base};"0"', "pct": "0.##%", "date": "dd-mmm-yyyy", "int": "0"}
    thin = Side(style="thin", color="A6A6A6")
    style = {"title": Font(size=13, bold=True), "line": Font(color="595959"), "bold": Font(bold=True),
             "foot": Font(size=9, color="808080"), "under": Border(bottom=thin), "over": Border(top=thin),
             "flag": PatternFill("solid", fgColor="FFF2CC"), "left": Alignment(horizontal="left"),
             "right": Alignment(horizontal="right")}
    wb = Workbook(write_only=True)
    who, when = getpass.getuser(), datetime.now().strftime("%d-%b-%Y %H:%M")
    used = set()
    for s in sheets:
        tab = "".join(c for c in (s.tab or s.title) if c not in "[]:*?/\\")[:31] or "Sheet"
        while tab.lower() in used:
            tab = tab[:28] + f" {len(used)}"
        used.add(tab.lower())
        ws = wb.create_sheet(tab)
        cols, n = list(s.frame.columns), len(s.frame)
        head = len(s.lines) + 3
        helper = _helper_columns(s) if s.chart and n else {}
        first = len(cols) + 20
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = f"A{head + 1}"
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        sample = s.frame.head(2000)
        for j, col in enumerate(cols, 1):
            longest = max([len(str(col))] + [len(_preview(v, s.kinds.get(col), units, decimals)) for v in sample[col]])
            ws.column_dimensions[get_column_letter(j)].width = min(60, max(longest + 3, 12))
        for k in range(len(helper)):
            ws.column_dimensions[get_column_letter(first + k)].hidden = True

        def cell(value, font=None, border=None, fill=None, align=None, number=None):
            c = WriteOnlyCell(ws, value=value)
            if font:
                c.font = font
            if border:
                c.border = border
            if fill:
                c.fill = fill
            if align:
                c.alignment = align
            if number:
                c.number_format = number
            return c

        ws.append([cell(s.title, style["title"])])
        for line in s.lines:
            ws.append([cell(line, style["line"])])
        ws.append([])
        header = [cell(str(col), style["bold"], style["under"],
                       align=style["left"] if s.kinds.get(col, "text") in ("text", "date") else style["right"])
                  for col in cols]
        if helper:
            header += [None] * (first - 1 - len(cols)) + list(helper)
        ws.append(header)
        numbers = [fmt.get(s.kinds.get(col)) for col in cols]
        aligns = [style["left"] if s.kinds.get(col) == "date" else None for col in cols]
        flags = set(s.flag)
        for i, row in enumerate(s.frame.itertuples(index=False)):
            last, hot = s.total and i == n - 1, i in flags
            out = []
            for v, number, align in zip(row, numbers, aligns):
                v = None if (v is None or (isinstance(v, float) and pd.isna(v))) else (v.item() if hasattr(v, "item") else v)
                if last or hot or number:
                    out.append(cell(v, style["bold"] if last else None, style["over"] if last else None,
                                    style["flag"] if hot and not last else None, align, number))
                else:
                    out.append(v)
            if helper:
                out += [None] * (first - 1 - len(cols)) + [series[i] for series in helper.values()]
            ws.append(out)
        ws.append([])
        ws.append([cell((f"Source: {s.source} · " if s.source else "") + f"exported {when} by {who} · "
                        f"USD, shown in {units}; cells hold full precision", style["foot"])])
        if helper:
            _chart(ws, s, head, units, fmt["money"], first, helper)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _helper_columns(s: Sheet) -> dict:
    c = s.chart
    ys, gs = list(s.frame[c["y"]]), [str(g) for g in s.frame[c["group"]]]
    n = len(ys)
    return {g: [float(ys[i]) if (gs[i] == g or (i + 1 < n and gs[i + 1] == g)) else None for i in range(n)]
            for g in dict.fromkeys(gs)}


def _preview(v, kind, units, decimals) -> str:
    if kind == "money" and isinstance(v, (int, float)) and not pd.isna(v):
        return f"{v / {'1': 1, 'k': 1e3, 'mn': 1e6, 'bn': 1e9}[units]:,.{decimals}f}"
    return str(v)


def _chart(ws, s: Sheet, head: int, units: str, money_fmt: str, first: int, helper: dict) -> None:
    from openpyxl.chart import LineChart, Reference, Series
    from openpyxl.utils import get_column_letter

    c, cols, n = s.chart, list(s.frame.columns), len(s.frame)
    ch = LineChart()
    ch.title, ch.height, ch.width = c.get("title"), 8, 20
    ch.display_blanks, ch.visible_cells_only = "gap", False
    ch.y_axis.title = f"USD {units}"
    ch.y_axis.number_format = money_fmt.split(";")[0]
    ch.y_axis.majorGridlines = None
    ch.x_axis.number_format = "dd-mmm"
    ch.x_axis.tickLblPos = "low"
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.legend.position = "b"
    vals = [float(v) for v in s.frame[c["y"]]]
    lo, hi = min(vals), max(vals)
    pad = (hi - lo) * 0.15 or abs(hi) * 0.05 or 1.0
    raw = (hi - lo + 2 * pad) / 5
    mag = 10 ** math.floor(math.log10(raw))
    step = mag * next(m for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    ch.y_axis.scaling.min = math.floor((lo - pad) / step) * step
    ch.y_axis.scaling.max = math.ceil((hi + pad) / step) * step
    ch.y_axis.majorUnit = step
    for k, g in enumerate(helper):
        series = Series(Reference(ws, min_col=first + k, min_row=head, max_row=head + n), title_from_data=True)
        colour = SCENARIO_COLOURS[k % len(SCENARIO_COLOURS)]
        series.graphicalProperties.line.solidFill = colour
        series.graphicalProperties.line.width = 22000
        series.marker.symbol, series.marker.size = "circle", 5
        series.marker.graphicalProperties.solidFill = colour
        series.marker.graphicalProperties.line.solidFill = colour
        series.smooth = False
        ch.series.append(series)
    ch.set_categories(Reference(ws, min_col=cols.index(c["x"]) + 1, min_row=head + 1, max_row=head + n))
    ch.anchor = f"{get_column_letter(len(cols) + 2)}{head}"
    ws.add_chart(ch)
